"""Text extraction from kit documents and the command-mention index (P1-R27).

Documents are untrusted: they are parsed as text only. PDF via pypdf, DOCX via python-docx (macros are never
executed), HTML via the stdlib parser. Failures become warnings, never crashes.
"""

import re
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import ClassVar

import docx
from pypdf import PdfReader


class DocumentError(Exception):
    """Text could not be extracted from a document."""


class _TextOnly(HTMLParser):
    _SKIP: ClassVar[frozenset[str]] = frozenset({"script", "style", "noscript", "template"})

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skipping = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._SKIP:
            self._skipping += 1
        elif tag in {"br", "p", "div", "li", "tr", "h1", "h2", "h3", "h4", "pre"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP and self._skipping:
            self._skipping -= 1

    def handle_data(self, data: str) -> None:
        if not self._skipping:
            self.parts.append(data)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            reader = PdfReader(path, strict=False)
            if reader.is_encrypted:
                raise DocumentError("PDF is encrypted")
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        if suffix == ".docx":
            document = docx.Document(str(path))
            lines = [p.text for p in document.paragraphs]
            for table in document.tables:
                lines += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
            return "\n".join(lines)
        if suffix in {".html", ".htm"}:
            parser = _TextOnly()
            parser.feed(path.read_text(encoding="utf-8", errors="replace"))
            return "".join(parser.parts)
        if suffix in {".md", ".txt"}:
            return path.read_text(encoding="utf-8", errors="replace")
    except DocumentError:
        raise
    except Exception as exc:  # malformed third-party documents must not crash intake
        raise DocumentError(f"{type(exc).__name__}: {exc}") from exc
    raise DocumentError(f"unsupported document type '{suffix}'")


# ---- command-mention index -------------------------------------------------------------------------------

_TOKEN = re.compile(r"\b[A-Z][A-Z0-9_]{2,}\b")
# Words that introduce a command name in user guides: "type DIMALL", "command: DIMALL", "run DIMALL".
_INTRO = re.compile(r"(?:command|type|enter|run)\s*[:\"'`]?\s*$", re.IGNORECASE)


@dataclass(frozen=True)
class Mention:
    document: str
    line: int
    snippet: str


def index_command_mentions(texts: dict[str, str], known: set[str]) -> dict[str, list[Mention]]:
    """Find likely command names in documents.

    A token counts when it is a known command (from `features[].commands`), contains an underscore
    (typical for plugin commands such as ACME_DIMALL), or follows words like "type" / "command:".
    """
    known_upper = {k.upper() for k in known}
    index: dict[str, list[Mention]] = {}
    for doc_name, text in texts.items():
        for line_no, line in enumerate(text.splitlines(), start=1):
            for match in _TOKEN.finditer(line):
                token = match.group(0)
                if token in known_upper or "_" in token or _INTRO.search(line[: match.start()]):
                    snippet = line.strip()[:160]
                    index.setdefault(token, []).append(Mention(doc_name, line_no, snippet))
    return dict(sorted(index.items()))
